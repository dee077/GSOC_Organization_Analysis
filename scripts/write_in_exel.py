import openpyxl
from openpyxl.utils import get_column_letter

def create_excel_workbook():
    """Create a new Excel workbook and return the workbook and sheet."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'GSoC Organizations Analysis 2024' 
    return workbook, sheet

def write_headers(sheet):
    """Write headers to the Excel sheet."""
    headers = ['Organization', 'Project_Ideas', 'Selected_Students', 'Technologies', 'Topics']
    
    # Write the headers to the first row
    for col_num, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_num)
        sheet[f'{col_letter}1'] = header

def sort_org_data(org_data):
    """Sort the organization data alphabetically by organization name."""
    return sorted(org_data, key=lambda item: item['Organization'])

def add_organization_data_to_sheet(sheet, sorted_org_data):
    """Add the organization data to the Excel sheet."""
    for row_num, org in enumerate(sorted_org_data, 2):
        sheet[f'A{row_num}'] = org['Organization']
        add_project_ideas(sheet, row_num, org)
        sheet[f'C{row_num}'] = org['Selected_Students']
        sheet[f'D{row_num}'] = org['Technologies']
        sheet[f'E{row_num}'] = org['Topics']

def add_project_ideas(sheet, row_num, org):
    """Add project ideas with a hyperlink to the sheet."""
    project_ideas_text = f"{org['Number_of_Ideas']} Ideas"
    project_ideas_link = org['Project_Ideas']
    sheet[f'B{row_num}'] = f'=HYPERLINK("{project_ideas_link}", "{project_ideas_text}")'

def save_workbook(workbook, year):
    """Save the Excel workbook to the specified file."""
    excel_file_path = f'data/{year}.xlsx'
    workbook.save(excel_file_path)
    print(f'Data has been stored in {excel_file_path}')

def write_data_to_excel(year, org_data):
    """Main function to write the organization data to an Excel file."""
    # Create a new workbook and sheet
    workbook, sheet = create_excel_workbook()

    # Write headers
    write_headers(sheet)

    # Sort organization data by the organization name
    sorted_org_data = sort_org_data(org_data)

    # Add organization data to the sheet
    add_organization_data_to_sheet(sheet, sorted_org_data)

    # Save the workbook
    save_workbook(workbook, year)
